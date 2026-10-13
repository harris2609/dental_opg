import io

import frappe
from docx import Document as DocxDocument
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from frappe.tests import IntegrationTestCase

from dental_opg.dental_opg_reporting.docx_export import TEMPLATE_PATH, build_docx


class TestDocxExport(IntegrationTestCase):
	def setUp(self):
		self.patient = frappe.get_doc(
			doctype="OPG Patient",
			patient_name="Docx Export Test",
			age=45,
			sex="Female",
			registration_no=frappe.generate_hash(length=8),
		).insert()
		self.report = frappe.get_doc(
			doctype="OPG Report",
			patient=self.patient.name,
			reported_by="Dr. Test Reporter",
			verified_by="Dr. Test Verifier",
		).insert()

	def test_template_file_exists_and_opens(self):
		DocxDocument(str(TEMPLATE_PATH))

	def test_build_docx_fills_populated_sections_and_removes_empty_ones(self):
		self.report.save_findings(
			section_findings=[
				{"section": "Condylar Morphology", "finding": "csf-normal"},
			],
			tooth_findings=[],
		)
		self.report.reload()

		content = build_docx(self.report)
		result_doc = DocxDocument(io.BytesIO(content))
		all_text = "\n".join(p.text for p in result_doc.paragraphs)

		self.assertIn("Condylar morphology appears to be normal", all_text)
		self.assertNotIn("STYLOID PROCESS", all_text)

	def test_build_docx_fills_body_table_tokens(self):
		"""The patient/report metadata table is body-level (unlike the
		footer), but result_doc.paragraphs alone doesn't reach table cell
		content either - table text has to be read via result_doc.tables.
		Nothing previously asserted these tokens actually get filled, only
		that the underlying replace function walks doc.tables at all."""
		self.report.referred_by = "Dr. Referring Doctor"
		self.report.clinical_diagnosis = "Routine screening"
		self.report.save()

		content = build_docx(self.report)
		result_doc = DocxDocument(io.BytesIO(content))

		table_text = ""
		for table in result_doc.tables:
			for row in table.rows:
				for cell in row.cells:
					table_text += cell.text + "\n"

		self.assertIn("Docx Export Test", table_text)
		self.assertIn("45", table_text)
		self.assertIn(self.patient.registration_no, table_text)
		self.assertIn("Dr. Referring Doctor", table_text)
		self.assertIn("Routine screening", table_text)
		self.assertNotIn("{{", table_text)

	def test_build_docx_fills_footer_tokens(self):
		"""The footer (Reported by / Verified by) is a separate docx part
		from the body - not reachable via result_doc.paragraphs alone. A
		prior version of this test only checked body paragraphs and missed
		that these tokens were going unfilled."""
		content = build_docx(self.report)
		result_doc = DocxDocument(io.BytesIO(content))
		footer = result_doc.sections[0].footer

		footer_text = "\n".join(p.text for p in footer.paragraphs)
		for table in footer.tables:
			for row in table.rows:
				for cell in row.cells:
					footer_text += "\n" + cell.text

		self.assertIn("Dr. Test Reporter", footer_text)
		self.assertIn("Dr. Test Verifier", footer_text)
		self.assertNotIn("{{", footer_text)

	def _export(self, section_findings=None, tooth_findings=None):
		self.report.save_findings(section_findings=section_findings or [], tooth_findings=tooth_findings or [])
		self.report.reload()
		return DocxDocument(io.BytesIO(build_docx(self.report)))

	def _all_runs(self, doc):
		footer = doc.sections[0].footer
		paragraphs = list(doc.paragraphs) + list(footer.paragraphs)
		for table in list(doc.tables) + list(footer.tables):
			for row in table.rows:
				for cell in row.cells:
					paragraphs.extend(cell.paragraphs)
		return [run for p in paragraphs for run in p.runs if run.text.strip()]

	def test_no_blank_paragraphs_after_findings_start(self):
		doc = self._export(section_findings=[{"section": "Condylar Morphology", "finding": "csf-normal"}])
		texts = [p.text for p in doc.paragraphs]
		start = next(i for i, t in enumerate(texts) if t.strip().startswith("PANORAMIC RADIOGRAPH SHOWS"))
		blank = [
			i
			for i, p in enumerate(doc.paragraphs[start + 1 :], start + 1)
			if not p.text.strip()
			and not p._p.findall(".//" + qn("w:drawing"))
			and p._p.find("./" + qn("w:pPr") + "/" + qn("w:sectPr")) is None
		]
		self.assertEqual(blank, [])

	def test_section_headings_get_space_before(self):
		doc = self._export(section_findings=[{"section": "Condylar Morphology", "finding": "csf-normal"}])
		heading = next(p for p in doc.paragraphs if p.text.strip() == "CONDYLAR MORPHOLOGY:")
		self.assertEqual(heading.paragraph_format.space_before, Pt(6))

	def test_every_run_is_black_times_new_roman(self):
		doc = self._export(
			section_findings=[{"section": "Condylar Morphology", "finding": "csf-normal"}],
			tooth_findings=[{"tooth_number": "16", "finding_name": "Caries", "description": "Coronally - a"}],
		)
		runs = self._all_runs(doc)
		self.assertTrue(runs)
		for run in runs:
			self.assertEqual(run.font.name, "Times New Roman", run.text)
			self.assertEqual(run.font.color.rgb, RGBColor(0, 0, 0), run.text)

	def test_multi_line_description_uses_line_breaks(self):
		doc = self._export(
			tooth_findings=[
				{"tooth_number": "16", "finding_name": "Caries", "description": "Coronally - a\nRadicularly - b"}
			]
		)
		line = next(p for p in doc.paragraphs if "irt 16" in p.text)
		self.assertEqual(line.text, "irt 16 - Coronally - a\nRadicularly - b")
		self.assertTrue(line._p.findall(".//" + qn("w:br")))

	def test_footer_cells_centered_without_fill(self):
		doc = self._export()
		table = doc.sections[0].footer.tables[0]
		for cell in table.rows[0].cells:
			self.assertIsNone(cell._tc.get_or_add_tcPr().find(qn("w:shd")))
			self.assertEqual(cell.vertical_alignment, WD_CELL_VERTICAL_ALIGNMENT.CENTER)
			for p in cell.paragraphs:
				self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.CENTER)
