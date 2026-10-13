import copy
import io
from pathlib import Path

import frappe
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

from dental_opg.dental_opg_reporting.image_utils import get_annotated_image_bytes
from dental_opg.dental_opg_reporting.report_builder import build_report_sections

TEMPLATE_PATH = Path(__file__).parent / "templates" / "opg_report_template.docx"

BULLET_SECTIONS = {"MAXILLARY_ARCH", "MANDIBULAR_ARCH"}

SECTION_HEADINGS = {
	"CONDYLAR_MORPHOLOGY": "CONDYLAR MORPHOLOGY:",
	"STYLOID_PROCESS": "STYLOID PROCESS:",
	"MAXILLARY_SINUS": "MAXILLARY SINUS:",
	"INFERIOR_BORDER_OF_MANDIBLE": "INFERIOR BORDER OF MANDIBLE",
	"MAXILLARY_ARCH": "MAXILLARY ARCH:",
	"MANDIBULAR_ARCH": "MANDIBULAR ARCH:",
	"RADIOLOGIC_IMPRESSION": "RADIOLOGIC IMPRESSION",
}

FONT_NAME = "Times New Roman"
BLACK = RGBColor(0, 0, 0)
FINDINGS_START = "PANORAMIC RADIOGRAPH SHOWS"


def build_docx(report) -> bytes:
	"""Fill the clinic's docx template with one OPG Report's data, return the
	resulting .docx file as bytes."""
	doc = Document(str(TEMPLATE_PATH))

	patient = frappe.get_doc("OPG Patient", report.patient) if report.patient else None
	_replace_tokens(doc, _metadata_replacements(report, patient))
	_insert_image(doc, report)
	_fill_sections(doc, report)
	_remove_blank_paragraphs(doc)
	_space_section_headings(doc)
	_restyle_footer(doc)
	_apply_black_times_new_roman(doc)

	buffer = io.BytesIO()
	doc.save(buffer)
	return buffer.getvalue()


def _metadata_replacements(report, patient) -> dict:
	return {
		"{{patient_name}}": (patient.patient_name if patient else "") or "",
		"{{age}}": str(patient.age) if patient and patient.age else "",
		"{{sex}}": (patient.sex or "")[:1] if patient else "",
		"{{registration_no}}": (patient.registration_no if patient else "") or "",
		"{{report_date}}": frappe.utils.formatdate(report.report_date) if report.report_date else "",
		"{{referred_by}}": report.referred_by or "",
		"{{radiographic_projection}}": report.radiographic_projection or "OPG",
		"{{clinical_diagnosis}}": report.clinical_diagnosis or "",
		"{{reported_by}}": report.reported_by or "",
		"{{verified_by}}": report.verified_by or "",
	}


def _iter_table_paragraphs(table):
	for row in table.rows:
		for cell in row.cells:
			yield from cell.paragraphs


def _iter_all_paragraphs(doc):
	"""Every paragraph a token could live in: body, body tables, and the
	footer (its own paragraphs plus the "Reported by"/"Verified by" table -
	the footer is a separate part from the body, not reachable via
	doc.paragraphs/doc.tables, which is exactly how {{reported_by}} and
	{{verified_by}} went unfilled despite two prior reviews passing."""
	yield from doc.paragraphs
	for table in doc.tables:
		yield from _iter_table_paragraphs(table)
	footer = doc.sections[0].footer
	yield from footer.paragraphs
	for table in footer.tables:
		yield from _iter_table_paragraphs(table)


def _replace_tokens(doc, replacements: dict):
	for p in _iter_all_paragraphs(doc):
		for run in p.runs:
			text = run.text
			for token, value in replacements.items():
				text = text.replace(token, value)
			run.text = text


def _insert_image(doc, report):
	marker = next((p for p in doc.paragraphs if p.text == "{{OPG_IMAGE}}"), None)
	if marker is None:
		return
	marker.text = ""
	content = get_annotated_image_bytes(report)
	if content is None:
		return
	marker.alignment = 1  # center
	run = marker.add_run()
	run.add_picture(io.BytesIO(content), width=Inches(6))


def _fill_sections(doc, report):
	sections_by_key = {s["name"].replace(" ", "_"): s["lines"] for s in build_report_sections(report)}

	for key, heading_text in SECTION_HEADINGS.items():
		marker_text = "{{FINDINGS:%s}}" % key
		marker = next((p for p in doc.paragraphs if p.text == marker_text), None)
		if marker is None:
			continue

		lines = sections_by_key.get(key)
		if not lines:
			heading = next((p for p in doc.paragraphs if p.text.strip() == heading_text), None)
			if heading is not None:
				heading._p.getparent().remove(heading._p)
			marker._p.getparent().remove(marker._p)
			continue

		style = marker.style
		marker.text = lines[0]
		last = marker
		for line in lines[1:]:
			last = _insert_paragraph_after(last, line, style)
			if key in BULLET_SECTIONS:
				_copy_numbering(marker, last)


def _insert_paragraph_after(reference_paragraph, text, style):
	new_p_elem = OxmlElement("w:p")
	reference_paragraph._p.addnext(new_p_elem)
	new_p = Paragraph(new_p_elem, reference_paragraph._parent)
	new_p.text = text
	new_p.style = style
	return new_p


def _copy_numbering(from_paragraph, to_paragraph):
	from_pPr = from_paragraph._p.pPr
	if from_pPr is None:
		return
	numPr = from_pPr.find(qn("w:numPr"))
	if numPr is None:
		return
	to_pPr = to_paragraph._p.get_or_add_pPr()
	to_pPr.append(copy.deepcopy(numPr))


def _keeps_layout(paragraph):
	p = paragraph._p
	return bool(p.findall(".//" + qn("w:drawing"))) or p.find("./" + qn("w:pPr") + "/" + qn("w:sectPr")) is not None


def _remove_blank_paragraphs(doc):
	"""The template carries hardcoded empty paragraphs between sections, which
	show up as big gaps in the export. Drop every empty paragraph after the
	findings heading; spacing comes from _space_section_headings instead."""
	paragraphs = doc.paragraphs
	start = next((i for i, p in enumerate(paragraphs) if p.text.strip().startswith(FINDINGS_START)), None)
	if start is None:
		return
	for paragraph in paragraphs[start + 1 :]:
		if not paragraph.text.strip() and not _keeps_layout(paragraph):
			paragraph._p.getparent().remove(paragraph._p)


def _space_section_headings(doc):
	headings = set(SECTION_HEADINGS.values())
	for paragraph in doc.paragraphs:
		if paragraph.text.strip() in headings:
			paragraph.paragraph_format.space_before = Pt(6)


def _restyle_footer(doc):
	"""Plain, centered two-column Reported by / Verified by footer: no colored
	fills (text goes black in _apply_black_times_new_roman)."""
	for table in doc.sections[0].footer.tables:
		table.alignment = WD_TABLE_ALIGNMENT.CENTER
		for row in table.rows:
			for cell in row.cells:
				tc_pr = cell._tc.get_or_add_tcPr()
				for shading in tc_pr.findall(qn("w:shd")):
					tc_pr.remove(shading)
				cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
				for paragraph in cell.paragraphs:
					paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _apply_black_times_new_roman(doc):
	"""Every character outside the letterhead header in black Times New Roman
	(the clinic saw blue text and mixed fonts). Multi-line finding text needs
	no special handling: python-docx turns "\\n" into <w:br/> on assignment."""
	normal = doc.styles["Normal"]
	normal.font.name = FONT_NAME
	normal.font.color.rgb = BLACK
	normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT_NAME)
	for paragraph in _iter_all_paragraphs(doc):
		for run in paragraph.runs:
			run.font.name = FONT_NAME
			run.font.color.rgb = BLACK
			run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT_NAME)
